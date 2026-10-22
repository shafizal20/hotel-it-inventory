
# backend/app/services/report_service.py

import csv
import io

from PIL import Image as PILImage
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def safe(value):
    """Uppercase text and protect CSV against formula injection."""
    text = "" if value is None else str(value).upper()

    if text[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + text

    return text


def _is_image_data(value):
    """Check whether a value is binary image data."""
    return isinstance(value, (bytes, bytearray, memoryview))


def _thumbnail(image_bytes, max_size=(100, 75)):
    """Convert image bytes to a PNG thumbnail."""
    if not image_bytes:
        return None

    try:
        with PILImage.open(io.BytesIO(bytes(image_bytes))) as original:
            image = original.convert("RGB")
            image.thumbnail(max_size)

            output = io.BytesIO()
            image.save(output, format="PNG")
            output.seek(0)
            return output
    except Exception:
        return None


def make_csv(name, cols, rows):
    """CSV cannot embed pictures; represent image columns as labels."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    writer.writerow([f"{name.upper()} REPORT"])
    writer.writerow([str(col).upper() for col in cols])

    for row in rows:
        output_row = []

        for value in row:
            if _is_image_data(value):
                output_row.append("IMAGE AVAILABLE" if value else "")
            else:
                output_row.append(safe(value))

        writer.writerow(output_row)

    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def make_xlsx(name, cols, rows):
    """Create an Excel workbook with embedded item pictures."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = name[:31].replace("/", "-").replace("\\", "-")

    # Report title
    sheet.append([f"{name.upper()} REPORT"])
    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=max(len(cols), 1),
    )

    sheet["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="1E293B")
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30

    # Column headings
    sheet.append([str(col).upper() for col in cols])

    for cell in sheet[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    sheet.freeze_panes = "A3"

    if cols:
        sheet.auto_filter.ref = (
            f"A2:{get_column_letter(len(cols))}{len(rows) + 2}"
        )

    # Keep streams alive until the workbook is saved.
    image_streams = []

    for row_index, row in enumerate(rows, start=3):
        sheet.row_dimensions[row_index].height = 60

        for col_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=col_index)

            if _is_image_data(value):
                # Never write raw image bytes into an Excel cell.
                thumbnail = _thumbnail(value, max_size=(100, 75))

                if thumbnail is not None:
                    picture = ExcelImage(thumbnail)
                    picture.width = 80
                    picture.height = 55
                    sheet.add_image(picture, cell.coordinate)
                    image_streams.append(thumbnail)
                else:
                    cell.value = "No image"
            else:
                cell.value = value

            cell.alignment = Alignment(
                vertical="center",
                horizontal="left",
                wrap_text=True,
            )

    # Set readable column widths.
    for column_index, column in enumerate(cols, start=1):
        heading = str(column).lower()

        if heading in ("pic", "picture", "image"):
            width = 16
        elif heading in (
            "item",
            "item_name",
            "location",
            "serial_number",
        ):
            width = 25
        else:
            width = min(max(len(str(column)) + 4, 12), 20)

        sheet.column_dimensions[
            get_column_letter(column_index)
        ].width = width

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def make_pdf(name, cols, rows):
    """Create a PDF report with image thumbnails."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Image as PDFImage,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buffer = io.BytesIO()
    styles = getSampleStyleSheet()

    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
    )

    header = [
        Paragraph(str(col).upper(), styles["BodyText"])
        for col in cols
    ]
    table_data = [header]
    image_streams = []

    for row in rows:
        output_row = []

        for value in row:
            if _is_image_data(value):
                thumbnail = _thumbnail(value, max_size=(160, 110))

                if thumbnail is not None:
                    thumbnail.seek(0)
                    image_streams.append(thumbnail)
                    output_row.append(
                        PDFImage(
                            thumbnail,
                            width=28 * mm,
                            height=19 * mm,
                        )
                    )
                else:
                    output_row.append("No image")
            else:
                text = "" if value is None else str(value)
                output_row.append(
                    Paragraph(
                        text[:250],
                        styles["BodyText"],
                    )
                )

        table_data.append(output_row)

    if not cols:
        return b""

    available_width = landscape(A4)[0] - 30 * mm
    col_widths = [available_width / len(cols)] * len(cols)

    for index, column in enumerate(cols):
        if str(column).lower() in ("pic", "picture", "image"):
            col_widths[index] = 32 * mm

    # Scale all columns proportionally to fit the page.
    total_width = sum(col_widths)
    if total_width > available_width:
        scale = available_width / total_width
        col_widths = [width * scale for width in col_widths]

    table = Table(
        table_data,
        colWidths=col_widths,
        repeatRows=1,
    )

    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        (
            "ROWBACKGROUNDS",
            (0, 1),
            (-1, -1),
            [colors.white, colors.HexColor("#f1f5f9")],
        ),
    ]))

    document.build([
        Paragraph(f"{name.upper()} REPORT", styles["Title"]),
        Spacer(1, 10),
        table,
    ])

    return buffer.getvalue()
